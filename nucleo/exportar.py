"""
Exporta la tabla de triplos a CSV y XLSX.
"""

import csv
import os


def exportar_csv(filas, catalogo, ruta_csv):
    """Genera un archivo CSV con la tabla del triplo."""
    with open(ruta_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        # Encabezado
        writer.writerow(["#", "Dato Objeto", "Dato Fuente", "Operador"])
        # Filas
        for fila in filas:
            writer.writerow([
                fila.numero,
                fila.objeto or "",
                fila.fuente or "",
                fila.operador or ""
            ])
        # Leyenda
        writer.writerow([])
        writer.writerow(["# LEYENDA DE TOKENS"])
        writer.writerow(["Token", "Lexema", "Categoría"])
        for id_sim, lexema, categoria in catalogo.leyenda():
            writer.writerow([id_sim, lexema, categoria])


def exportar_excel(filas, catalogo, ruta_xlsx):
    """Genera un archivo XLSX con formato."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise ImportError(
            "Se requiere la librería 'openpyxl' para exportar a Excel.\n"
            "Instálala con: pip install openpyxl"
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "Triplo"

    # ---------- Estilos ----------
    borde = Border(
        left=Side(style="thin", color="cccccc"),
        right=Side(style="thin", color="cccccc"),
        top=Side(style="thin", color="cccccc"),
        bottom=Side(style="thin", color="cccccc"),
    )
    centrado = Alignment(horizontal="center", vertical="center")

    fuente_encabezado = Font(bold=True, color="FFFFFF", size=11)
    fondo_encabezado = PatternFill("solid", fgColor="2C3E50")

    colores_por_tipo = {
        "asignacion":     PatternFill("solid", fgColor="D6EAF8"),   # azul claro
        "aritmetica":     PatternFill("solid", fgColor="FEF9E7"),   # amarillo claro
        "relacional":     PatternFill("solid", fgColor="D5F4E6"),   # verde claro
        "logico":         PatternFill("solid", fgColor="FDEBD0"),   # naranja claro
        "salto_cond":     PatternFill("solid", fgColor="FADBD8"),   # rojo claro
        "salto_incond":   PatternFill("solid", fgColor="E5E7E9"),   # gris claro
        "otro":           PatternFill("solid", fgColor="FFFFFF"),
    }

    # ---------- Encabezado ----------
    encabezados = ["#", "Dato Objeto", "Dato Fuente", "Operador"]
    for col, texto in enumerate(encabezados, start=1):
        celda = ws.cell(row=1, column=col, value=texto)
        celda.font = fuente_encabezado
        celda.fill = fondo_encabezado
        celda.alignment = centrado
        celda.border = borde

    # ---------- Filas ----------
    for idx, fila in enumerate(filas, start=2):
        valores = [
            fila.numero,
            fila.objeto or "",
            fila.fuente or "",
            fila.operador or ""
        ]
        color = colores_por_tipo.get(fila.tipo_fila, colores_por_tipo["otro"])
        for col, valor in enumerate(valores, start=1):
            celda = ws.cell(row=idx, column=col, value=valor)
            celda.fill = color
            celda.border = borde
            celda.alignment = centrado

    # ---------- Anchos de columna ----------
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 14

    # ---------- Hoja 2: Leyenda ----------
    ws2 = wb.create_sheet("Leyenda")
    encabezados2 = ["Token", "Lexema", "Categoría"]
    for col, texto in enumerate(encabezados2, start=1):
        celda = ws2.cell(row=1, column=col, value=texto)
        celda.font = fuente_encabezado
        celda.fill = fondo_encabezado
        celda.alignment = centrado
        celda.border = borde

    for idx, (id_sim, lexema, categoria) in enumerate(catalogo.leyenda(), start=2):
        for col, valor in enumerate([id_sim, lexema, categoria], start=1):
            celda = ws2.cell(row=idx, column=col, value=valor)
            celda.border = borde
            celda.alignment = centrado

    ws2.column_dimensions["A"].width = 12
    ws2.column_dimensions["B"].width = 20
    ws2.column_dimensions["C"].width = 25

    wb.save(ruta_xlsx)


def exportar_ambos(filas, catalogo, ruta_base):
    """
    Exporta a CSV y XLSX usando el mismo nombre base.
    Devuelve (ruta_csv, ruta_xlsx).
    """
    # Quitar extensión si el usuario la puso
    if ruta_base.lower().endswith(".csv"):
        ruta_base = ruta_base[:-4]
    elif ruta_base.lower().endswith(".xlsx"):
        ruta_base = ruta_base[:-5]

    ruta_csv = ruta_base + ".csv"
    ruta_xlsx = ruta_base + ".xlsx"

    exportar_csv(filas, catalogo, ruta_csv)
    exportar_excel(filas, catalogo, ruta_xlsx)

    return ruta_csv, ruta_xlsx